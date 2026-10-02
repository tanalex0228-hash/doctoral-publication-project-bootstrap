import csv
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, render
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config.pagination import paginate_queryset
from doctoral_students.models import DoctoralStudentProfile
from publications.permissions import is_staff_actor
from .forms import StatisticsFilterForm
from .services import (
    STUDENT_SORTS,
    approved_publications_for_student,
    csv_safe_cell,
    export_ready_publications,
    export_ready_rows,
    filtered_approved_publications,
    statistics_overview,
    student_statistics_table,
    student_summary,
)
from .secretary_exports import CONFERENCE_HEADERS, JOURNAL_HEADERS, conference_rows, journal_rows


def _require_statistics_role(request):
    if not is_staff_actor(request.user):
        raise Http404("Statistics resource not found.")


def _filtered_records(request):
    form = StatisticsFilterForm(request.GET or None)
    if not form.is_valid():
        return form, filtered_approved_publications()
    return form, filtered_approved_publications(**form.cleaned_data)


PUBLICATION_EXPORT_HEADERS = [
    "學生", "學號", "指導教授", "成果標題", "成果類型", "Publication Index", "研究領域", "作者",
    "期刊／會議", "發表日期", "接受日期", "DOI", "ISSN", "語言", "核准日期", "Visibility", "發佈狀態",
]


def _xlsx_response(*, filename, worksheet_title, headers, rows):
    """Build a safe Excel download without exposing data outside the approved query."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = worksheet_title[:31]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    for row in rows:
        worksheet.append([csv_safe_cell(value) for value in row])
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        width = min(max((len(str(cell.value or "")) for cell in column), default=0) + 2, 40)
        worksheet.column_dimensions[get_column_letter(index)].width = width
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _valid_filtered_records(request):
    form, records = _filtered_records(request)
    return records if form.is_valid() else records.none()


def _student_statistics_rows(records, *, sort):
    for student in student_statistics_table(records, sort=sort):
        yield [
            student.display_name, student.student_number, student.approved_total, student.journal_total,
            student.conference_total, student.first_author_total, student.corresponding_author_total,
            student.advisor_coauthored_total, student.publication_index_summary or "",
            student.advisor_names or "", student.latest_publication_date or "", student.first_publication_type or "",
        ]


@login_required
def statistics_dashboard(request):
    _require_statistics_role(request)
    form, records = _filtered_records(request)
    return render(request, "statistics/dashboard.html", {
        "filter_form": form,
        "overview": statistics_overview(records),
    })


@login_required
def approved_publication_list(request):
    """Staff drill-down for the approved-only dashboard total."""
    _require_statistics_role(request)
    form, records = _filtered_records(request)
    publications = export_ready_publications(records)
    page_obj, pagination_query = paginate_queryset(request, publications)
    return render(request, "statistics/approved_list.html", {
        "filter_form": form,
        "publications": page_obj,
        "page_obj": page_obj,
        "pagination_query": pagination_query,
    })


@login_required
def approved_xlsx_export(request):
    _require_statistics_role(request)
    return _xlsx_response(
        filename="approved-publications.xlsx", worksheet_title="已核准成果",
        headers=PUBLICATION_EXPORT_HEADERS, rows=export_ready_rows(_valid_filtered_records(request)),
    )


@login_required
def student_statistics(request):
    _require_statistics_role(request)
    form, records = _filtered_records(request)
    requested_sort = request.GET.get("sort", "name")
    sort = requested_sort if requested_sort in STUDENT_SORTS else "name"
    students = student_statistics_table(records, sort=sort)
    page_obj, pagination_query = paginate_queryset(request, students)
    return render(request, "statistics/student_list.html", {
        "filter_form": form,
        "students": page_obj,
        "page_obj": page_obj,
        "pagination_query": pagination_query,
        "sort": sort,
    })


@login_required
def student_statistics_xlsx_export(request):
    _require_statistics_role(request)
    requested_sort = request.GET.get("sort", "name")
    sort = requested_sort if requested_sort in STUDENT_SORTS else "name"
    return _xlsx_response(
        filename="approved-student-statistics.xlsx", worksheet_title="學生正式成果統計",
        headers=["學生", "學號", "正式數", "Journal", "Conference", "第一作者", "通訊作者", "與指導教授合著", "Index 概況", "指導教授", "最近發表", "成果類型"],
        rows=_student_statistics_rows(_valid_filtered_records(request), sort=sort),
    )


@login_required
def student_statistics_detail(request, student_id):
    _require_statistics_role(request)
    student = get_object_or_404(DoctoralStudentProfile, pk=student_id)
    publications = approved_publications_for_student(student)
    return render(request, "statistics/student_detail.html", {
        "student": student,
        "publications": publications,
        "summary": student_summary(student),
    })


@login_required
def student_statistics_detail_xlsx_export(request, student_id):
    _require_statistics_role(request)
    student = get_object_or_404(DoctoralStudentProfile, pk=student_id)
    return _xlsx_response(
        filename=f"approved-publications-{student.student_number}.xlsx",
        worksheet_title=f"{student.display_name}成果",
        headers=PUBLICATION_EXPORT_HEADERS,
        rows=export_ready_rows(approved_publications_for_student(student)),
    )


@login_required
def export_ready_data(request):
    _require_statistics_role(request)
    form, records = _filtered_records(request)
    publications = export_ready_publications(records)
    page_obj, pagination_query = paginate_queryset(request, publications)
    return render(request, "statistics/export_ready.html", {
        "filter_form": form,
        "publications": page_obj,
        "page_obj": page_obj,
        "pagination_query": pagination_query,
    })


@login_required
def export_ready_xlsx_export(request):
    _require_statistics_role(request)
    return _xlsx_response(
        filename="export-ready-approved-publications.xlsx", worksheet_title="匯出準備資料",
        headers=PUBLICATION_EXPORT_HEADERS, rows=export_ready_rows(_valid_filtered_records(request)),
    )


@login_required
def csv_export(request):
    _require_statistics_role(request)
    form, records = _filtered_records(request)
    if not form.is_valid():
        records = records.none()
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="approved-publications.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow([
        "Student", "Student ID", "Advisors", "Publication Title", "Publication Type",
        "Publication Index", "Research Field", "Authors", "Journal / Conference",
        "Publication Date", "Acceptance Date", "DOI", "ISSN", "Language", "Approved Date",
        "Visibility", "Publishing State",
    ])
    for row in export_ready_rows(records):
        writer.writerow([csv_safe_cell(value) for value in row])
    return response


def _secretary_csv_response(filename, headers, rows):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([csv_safe_cell(value) for value in row])
    return response


@login_required
def secretary_journal_export(request):
    _require_statistics_role(request)
    return _secretary_csv_response("secretary-journal-publications.csv", JOURNAL_HEADERS, journal_rows())


@login_required
def secretary_conference_export(request):
    _require_statistics_role(request)
    return _secretary_csv_response("secretary-conference-publications.csv", CONFERENCE_HEADERS, conference_rows())
