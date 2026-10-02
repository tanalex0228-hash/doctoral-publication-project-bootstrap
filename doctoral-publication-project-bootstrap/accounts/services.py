from dataclasses import dataclass
from io import BytesIO

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook

from audit.services import record_event
from doctoral_students.models import DoctoralStudentProfile
from professors.models import Professor

from .models import Role, User, UserRole


IMPORT_HEADERS = (
    "序號", "姓名", "帳號", "學號", "身份", "單位", "預設密碼", "管理介面", "密碼登入", "權限群組", "入學年度",
)
ROLE_MAP = {"teacher": "advisor", "student": "student", "admin": "admin"}
GROUP_INPUT_MAP = {"teacher": "advisor", "advisor": "advisor", "student": "student", "admin": "admin"}
TRUE_VALUES = {"v", "y", "yes", "true", "1"}
FALSE_VALUES = {"x", "n", "no", "false", "0", ""}


@dataclass(frozen=True)
class ImportedAccount:
    row_number: int
    name: str
    email: str
    student_number: str
    identity: str
    unit: str
    default_password: str
    admin_interface: bool
    password_login: bool
    role_slug: str
    admission_year: int | None


def can_manage_accounts(user):
    return bool(
        user and user.is_authenticated and user.is_active
        and (user.is_superuser or user.has_project_role("admin"))
    )


def _value(value):
    return "" if value is None else str(value).strip()


def _flag(value, *, label, row_number):
    normalized = _value(value).lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise ValidationError(f"第 {row_number} 列「{label}」必須為 v/y 或 x/n。")


def _optional_year(value, *, row_number):
    value = _value(value)
    if not value:
        return None
    try:
        year = int(value)
    except (TypeError, ValueError):
        raise ValidationError(f"第 {row_number} 列「入學年度」必須為正整數或留白。")
    if year <= 0:
        raise ValidationError(f"第 {row_number} 列「入學年度」必須為正整數或留白。")
    return year


def _existing_email_query(emails):
    query = Q()
    for email in emails:
        query |= Q(email__iexact=email)
    return query


def parse_account_workbook(upload):
    """Validate the complete workbook before any account is written."""
    try:
        workbook = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("無法讀取 Excel 檔案，請確認檔案未損毀且為 .xlsx 格式。") from exc

    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = tuple(_value(value) for value in next(rows, ()))
    if headers != IMPORT_HEADERS:
        raise ValidationError("欄位順序或名稱不符。必須完全符合：" + "、".join(IMPORT_HEADERS))

    accounts = []
    errors = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and _value(value) for value in row):
            continue
        if len(row) != len(IMPORT_HEADERS):
            errors.append(f"第 {row_number} 列欄位數不正確。")
            continue
        try:
            name, email, student_number, identity, unit, password, admin_ui, password_login, role, year = (
                _value(row[index]) for index in range(1, len(IMPORT_HEADERS))
            )
            identity = identity.lower()
            role = GROUP_INPUT_MAP.get(role.lower(), "")
            if not name:
                raise ValidationError(f"第 {row_number} 列「姓名」不可空白。")
            if not email or "@" not in email:
                raise ValidationError(f"第 {row_number} 列「帳號」必須為有效電子郵件。")
            if identity not in ROLE_MAP:
                raise ValidationError(f"第 {row_number} 列「身份」必須為 teacher、student 或 admin。")
            if not role:
                raise ValidationError(f"第 {row_number} 列「權限群組」必須為 teacher、student 或 admin。")
            if role != ROLE_MAP[identity]:
                raise ValidationError(f"第 {row_number} 列「身份」與「權限群組」必須相符。")
            if identity == "student" and not student_number:
                raise ValidationError(f"第 {row_number} 列學生必須填寫「學號」。")
            can_password_login = _flag(password_login, label="密碼登入", row_number=row_number)
            if can_password_login and not password:
                raise ValidationError(f"第 {row_number} 列啟用密碼登入時必須填寫「預設密碼」。")
            accounts.append(ImportedAccount(
                row_number=row_number, name=name, email=email, student_number=student_number,
                identity=identity, unit=unit, default_password=password,
                admin_interface=_flag(admin_ui, label="管理介面", row_number=row_number),
                password_login=can_password_login, role_slug=role,
                admission_year=_optional_year(year, row_number=row_number),
            ))
        except ValidationError as exc:
            errors.extend(exc.messages)

    if not accounts and not errors:
        errors.append("Excel 檔案沒有可匯入的資料列。")
    for field, label in (("email", "帳號"), ("name", "姓名"), ("student_number", "學號")):
        seen = {}
        for account in accounts:
            value = getattr(account, field)
            if not value:
                continue
            key = value.lower() if field == "email" else value
            if key in seen:
                errors.append(f"第 {account.row_number} 列「{label}」與第 {seen[key]} 列重複。")
            seen[key] = account.row_number
    if errors:
        raise ValidationError(errors)

    emails = [account.email for account in accounts]
    names = [account.name for account in accounts]
    student_numbers = [account.student_number for account in accounts if account.student_number]
    existing_errors = []
    if User.objects.filter(_existing_email_query(emails)).exists():
        existing_errors.append("檔案中有帳號已存在。")
    if User.objects.filter(username__in=names).exists():
        existing_errors.append("檔案中有姓名已作為既有使用者名稱。")
    if student_numbers and DoctoralStudentProfile.objects.filter(student_number__in=student_numbers).exists():
        existing_errors.append("檔案中有學號已存在。")
    missing_roles = set(ROLE_MAP.values()) - set(Role.objects.filter(slug__in=ROLE_MAP.values(), is_active=True).values_list("slug", flat=True))
    if missing_roles:
        existing_errors.append("系統缺少啟用的權限群組：" + "、".join(sorted(missing_roles)))
    if existing_errors:
        raise ValidationError(existing_errors)
    return accounts


@transaction.atomic
def import_accounts(*, actor, accounts, request_id=None):
    if not can_manage_accounts(actor):
        raise PermissionDenied("僅限系統管理員批量建立帳號。")
    roles = Role.objects.in_bulk(ROLE_MAP.values(), field_name="slug")
    created = []
    for account in accounts:
        user = User(username=account.name, first_name=account.name, email=account.email, is_staff=account.admin_interface)
        if account.password_login:
            user.set_password(account.default_password)
        else:
            user.set_unusable_password()
        user.save()
        UserRole.objects.create(user=user, role=roles[account.role_slug], assigned_by=actor)
        if account.identity == "student":
            DoctoralStudentProfile.objects.create(
                user=user, student_number=account.student_number, display_name=account.name,
                admission_year=account.admission_year,
            )
        elif account.identity == "teacher":
            Professor.objects.create(user=user, display_name=account.name, email=account.email)
        record_event(
            actor=actor, action="account.bulk_imported", target=user, request_id=request_id,
            metadata={"identity": account.identity, "role": account.role_slug, "unit": account.unit,
                      "admin_interface": account.admin_interface, "password_login": account.password_login,
                      "admission_year": account.admission_year},
        )
        created.append(user)
    return created


@transaction.atomic
def replace_user_role(*, actor, user, role_slug, request_id=None):
    if not can_manage_accounts(actor):
        raise PermissionDenied("僅限系統管理員變更權限群組。")
    role = Role.objects.get(slug=role_slug, is_active=True)
    user = User.objects.select_for_update().get(pk=user.pk)
    previous = list(user.user_roles.values_list("role__slug", flat=True))
    UserRole.objects.filter(user=user).delete()
    UserRole.objects.create(user=user, role=role, assigned_by=actor)
    record_event(actor=actor, action="account.role_replaced", target=user, request_id=request_id,
                 metadata={"previous_roles": previous, "role": role.slug})
    return user
